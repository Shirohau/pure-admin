#!/usr/bin/env python
# -*- coding: UTF-8 -*-
"""
@Project ：backend
@File    ：sync_export_mixin.py
@Author  ：李小涛
@Date    ：2026/8/3
@Explain : 同步导出混入（基于 django-import-export）
          独立管理导出相关逻辑：
          文件下载、字段权限计算、同步导出、可导出字段查询。
"""
import codecs
import re
from datetime import datetime
from typing import Union

from django.core.exceptions import FieldDoesNotExist
from django.db.models import Max
from django.http import HttpResponse
from django.template.defaultfilters import capfirst
from django.utils.encoding import escape_uri_path
from drf_spectacular.utils import extend_schema
from import_export.formats import base_formats
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from tablib import Dataset

from apps.system.models.menu_field.models import MenuFieldModel
from apps.system.models.role_menu_field.models import RoleMenuFieldModel
from extends.drf.response import error_response, success_response

# 导出/下载链路所校验的功能权限 key（可配置扩展点）：
# 只有配置在此的功能权限才参与导出字段过滤，判定为"全部满足才允许导出该字段"。
# 功能权限定义见 FUNC_PERMISSION_DEFINITIONS（如 can_download 可下载 / can_print 可打印）。
# 当前系统打印为纯前端能力（无后端打印接口），故仅 can_download 接入导出/下载校验；
# 后续若有功能权限（如打印）也需接入导出链路，只需把对应 key 追加到此元组，无需改动判定逻辑。
EXPORT_FUNC_PERMISSION_KEYS = ("can_download",)

# 导出/导入格式名称 → 格式类映射（集中管理，便于扩展）
SUPPORTED_EXPORT_FORMATS = {
    "csv": base_formats.CSV,
    "xlsx": base_formats.XLSX,
    # 可扩展："json": base_formats.JSON, ...
}


def make_disposition(filename: str) -> str:
    """
    生成兼容中文文件名的 Content-Disposition 响应头

    实现说明：
    同时提供 ASCII 兜底文件名（filename）与 RFC 5987 编码文件名（filename*），
    保证旧客户端（只识别 ASCII）与新客户端（支持 UTF-8）都能正确下载中文文件。

    Args:
        filename: 原始文件名（可含中文）

    Returns:
        Content-Disposition 响应头字符串
    """

    def to_ascii(name: str) -> str:
        # 将非安全字符替换为下划线，生成纯 ASCII 文件名
        return re.sub(r"[^a-zA-Z0-9._-]", "_", name)

    ascii_name = to_ascii(filename)
    encoded_name = escape_uri_path(filename)
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{encoded_name}'


def file_download_response(
        data: Union[bytes, Dataset],
        file_format: str,
        filename: str,
        column_widths: dict = None,
) -> HttpResponse:
    """
    构造标准的文件下载 HttpResponse，自动处理中文文件名

    CSV 响应在最终输出处确保包含 UTF-8 BOM，避免格式类因模块导入顺序
    被提前缓存时退回原始 CSV 实现。

    Args:
        data: 文件内容（Dataset 或 bytes）
        file_format: 文件格式字符串，如 "csv"、"xlsx"
        filename: 下载文件名（可含中文）
        column_widths: 列宽配置字典，格式为 {"列名": 宽度}，仅对 xlsx 格式生效

    Returns:
        文件下载响应（含 Content-Disposition 头）
    """
    format_class = SUPPORTED_EXPORT_FORMATS.get(file_format)()
    content_type = format_class.get_content_type()
    export_data = format_class.export_data(data)

    # xlsx 格式：根据 column_widths 设置列宽
    if file_format == "xlsx" and column_widths:
        from io import BytesIO
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        # 加载已导出的工作簿
        wb = load_workbook(BytesIO(export_data))
        ws = wb.active

        # 获取表头行，用于匹配列名到列索引
        headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []

        # 遍历列宽配置，设置对应列的宽度
        for col_idx, header in enumerate(headers, start=1):
            if header in column_widths:
                ws.column_dimensions[get_column_letter(col_idx)].width = column_widths[header]

        # 重新序列化工作簿
        output = BytesIO()
        wb.save(output)
        export_data = output.getvalue()

    if file_format == "csv":
        if isinstance(export_data, str):
            export_data = export_data.encode("utf-8")
        if not export_data.startswith(codecs.BOM_UTF8):
            export_data = codecs.BOM_UTF8 + export_data

    response = HttpResponse(export_data, content_type=content_type)
    response["Content-Disposition"] = make_disposition(filename)
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response


def _get_permitted_export_field_attrs(model_class, user):
    """
    根据当前用户的字段级权限，返回"字段是否允许导出"的判定函数

    权限判定规则（与前端 fieldAuthManager.ts、RoleMenuFieldModel 语义对齐）：
      1. 非模型字段（自定义导出列，如 creator/updater）→ 不受字段权限管理，默认放行；
      2. 模型字段未被 MenuFieldModel 管理 → 不受字段权限管理，默认放行；
      3. 模型字段受管理但用户角色未显式分配（无 RoleMenuFieldModel 记录）
         → 遵循"未显式分配权限，则默认拥有全部权限"，放行；
      4. 模型字段受管理且已显式分配 → 遵循"一旦分配了权限，则仅拥有被分配的权限"，
         需同时满足：
           - permission_level > 0（非"禁止访问"，禁止字段不可见也不可导出）
           - func_permissions 中 can_download = True（可下载为独立功能权限，仅对导出/下载链路生效）

    多角色合并策略（用户可拥有多个角色，取最宽松）：
      - permission_level 取最大值（Max）；
      - 导出功能权限（EXPORT_FUNC_PERMISSION_KEYS，如 can_download）取任一角色允许即允许（Max of 0/1）。

    Args:
        model_class: 导出资源对应的 Django 模型类
        user: 当前请求用户（需含 role 关联）

    Returns:
        可调用对象 is_permitted(attr) -> bool；返回 None 表示不限制（无角色体系等）。
    """
    if not hasattr(user, "role"):
        return None  # 无角色体系，不限制

    role_ids = list(user.role.values_list("id", flat=True))
    if not role_ids:
        return None

    model_label = model_class._meta.label_lower

    # 模型自身字段名（用于区分自定义导出列与真实模型字段）
    model_field_names = set(f.name for f in model_class._meta.get_fields())

    # 查询哪些字段被 MenuFieldModel 管理（未管理的字段不受字段权限体系约束）
    managed_fields = set(
        MenuFieldModel.objects.filter(model=model_label).values_list("field_name", flat=True)
    )

    # 查询当前用户对这些受管字段的显式分配记录（多角色取最宽松：等级取 Max、功能权限取任一允许）
    # 动态 annotate：为每个导出功能权限 key 生成 max_<key>=Max("func_permissions__<key>")
    # （func_permissions 为 JSON 字段，按 key 聚合；新增功能权限无需改动此处）
    annotate_kwargs = {
        f"max_{key}": Max(f"func_permissions__{key}")
        for key in EXPORT_FUNC_PERMISSION_KEYS
    }
    explicit_perms = RoleMenuFieldModel.objects.filter(
        role_id__in=role_ids,
        menu_field__model=model_label
    ).values(
        "menu_field__field_name"
    ).annotate(
        max_level=Max("permission_level"),
        **annotate_kwargs,
    )

    # 字段名 → 合并后的显式权限（{等级, 各功能权限是否允许}）
    field_perms = {
        item["menu_field__field_name"]: {
            "level": item["max_level"],
            **{key: bool(item[f"max_{key}"]) for key in EXPORT_FUNC_PERMISSION_KEYS},
        }
        for item in explicit_perms
    }

    def is_permitted(attr):
        """判断单个字段（attribute）是否允许导出"""
        # 自定义导出列（不在模型中）→ 不受字段权限管理，放行
        if attr not in model_field_names:
            return True

        # 模型字段：未纳入权限体系 → 放行
        if attr not in managed_fields:
            return True

        # 受管字段：未显式分配 → 默认拥有全部权限，放行
        perm = field_perms.get(attr)
        if perm is None:
            return True

        # 已显式分配 → 仅拥有被分配的权限：非禁止访问（NULL=未显式分配数据等级，视为非禁止）
        # + 导出功能权限全部允许，才允许导出
        return perm["level"] != 0 and all(
            perm[key] for key in EXPORT_FUNC_PERMISSION_KEYS
        )

    return is_permitted


def _compute_permitted_export_attrs(resource_class, user):
    """
    根据当前用户的字段级权限，计算允许导出的字段 attribute 列表

    供同步导出（sync_export）与异步导出（get_export_resource_kwargs）共用，
    保证两条链路的字段权限规则完全一致。

    实现说明：
    - 返回 None 表示不限制（无角色体系或未分配任何角色）；
    - 必须用列表而非集合：该结果可能存入 ExportJob.resource_kwargs（JSONField），
      set 无法被 json.dumps 序列化（TypeError: Object of type set is not JSON serializable）。

    Args:
        resource_class: 导出资源类（import_export.Resource 子类）
        user: 当前请求用户

    Returns:
        允许导出的字段 attribute 列表；None 表示不限制
    """
    is_permitted = _get_permitted_export_field_attrs(
        model_class=resource_class._meta.model,
        user=user,
    )
    if is_permitted is None:
        return None

    all_fields = resource_class().get_export_fields()
    return [
        f.attribute for f in all_fields
        if not f.attribute or is_permitted(f.attribute)
    ]


class ExportMixin:
    """
    同步导出混入：
        - sync_export: 同步导出数据（超过最大行数时提示改用异步导出）
        - export_fields: 获取当前用户有权导出的字段列表

    可配置属性（子类需按需配置）：
        - export_resource_class: 导出资源类（必配，未配置时断言报错）
        - export_max_rows: 同步导出最大行数限制，默认 500
    """

    export_resource_class = None  # 同步导出资源类
    export_max_rows = 500  # 同步导出最大行数

    def get_resource_class(self, prefix):
        """
        根据动作前缀获取对应的资源类

        Args:
            prefix: 动作前缀，如 "export" / "import"

        Returns:
            资源类（未配置时抛出断言异常提示）
        """
        action_resource_name = f"{prefix}_resource_class"
        action_resource_class = getattr(self, action_resource_name, None)
        assert action_resource_class, f"{self.__class__.__name__} 请配置对应的 {prefix}_resource_class"
        return action_resource_class

    @extend_schema(summary="导出", extensions={"x-function": "SyncExport"})
    @action(methods=["get"], detail=False)
    def sync_export(self, request, *args, **kwargs):
        """
        同步导出数据到文件

        请求参数：
            - selected_fields: 指定导出字段，如 ?selected_fields=name,age（可选）
            - file_format: 文件格式，如 csv / xlsx，默认 csv

        导出字段受字段级权限（可下载权限）约束，规则与 export_fields 接口一致：
            - 未显式分配权限的字段默认拥有全部权限（可导出）；
            - 已显式分配的字段，仅当"非禁止访问且可下载"时才可导出。

        Args:
            request: 请求对象

        Returns:
            文件下载响应；行数超限时返回错误响应
        """
        # 解析导出字段参数（逗号分隔 → 列表，未传则导出全部字段）
        selected_fields = request.query_params.get("selected_fields")
        selected_fields = selected_fields.split(",") if selected_fields else None
        # 获取文件格式并转为小写，默认 csv
        file_format = request.query_params.get("file_format", "csv").lower()
        # 边界处理：白名单校验，非法格式直接返回错误提示，避免格式类为 None 时抛 TypeError
        if file_format not in SUPPORTED_EXPORT_FORMATS:
            return error_response(
                message=f"不支持的文件格式：{file_format}，仅支持 {', '.join(SUPPORTED_EXPORT_FORMATS.keys())}"
            )
        # 获取过滤后的数据集（含数据权限、搜索等过滤器）
        queryset = self.filter_queryset(self.get_queryset())

        # 根据当前用户的字段级权限，计算允许导出的字段 attribute 集合
        # （双保险：即使前端未传 selected_fields，资源层也只导出有下载权限的字段）
        resource_class = self.get_resource_class("export")
        permitted_attrs = _compute_permitted_export_attrs(
            resource_class=resource_class,
            user=request.user,
        )

        # 实例化资源类，传入字段选择与下载权限过滤集合
        resource = resource_class(_selected_fields=selected_fields, _permitted_attrs=permitted_attrs)
        # 执行数据导出
        export_data = resource.export(queryset)
        if len(export_data) > self.export_max_rows:
            return error_response(
                message=f"文件行数超过{self.export_max_rows}行，请使用后台导出",
                data={"max_rows": self.export_max_rows, "actual_rows": len(export_data)},
            )

        # 构造导出文件名：模型名_导出数据_时间戳.格式
        model_name = capfirst(queryset.model._meta.verbose_name.replace(" ", ""))
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        filename = f"{model_name}_导出数据_{timestamp}.{file_format}"

        # 返回文件下载响应
        return file_download_response(data=export_data, file_format=file_format, filename=filename)

    @extend_schema(summary="可导出字段", extensions={"x-function": "ExportFields", "x-assign_permission": False})
    @action(methods=["get"], detail=False, permission_classes=[IsAuthenticated])
    def export_fields(self, request, *args, **kwargs):
        """
        获取当前用户有权导出的字段列表（权限过滤在接口层完成）

        权限规则（与 sync_export 导出链路保持一致）：
            - 未显式分配权限的字段默认拥有全部权限（可导出）；
            - 已显式分配的字段，仅当"非禁止访问且可下载"时才可导出。

        Args:
            request: 请求对象

        Returns:
            统一成功响应，data 为 [{value, label}, ...] 字段列表
        """
        resource = self.get_resource_class("export")()
        # 获取所有原始导出字段（未过滤）
        all_fields = resource.get_export_fields()
        # 获取字段级下载权限判定函数（None 表示不限制）
        is_permitted = _get_permitted_export_field_attrs(
            model_class=resource._meta.model,
            user=request.user,
        )

        result = []
        for field in all_fields:
            attr = field.attribute  # 模型字段名，如 "name"
            column_name = field.column_name  # 导出列名
            # 关键：在此处应用权限过滤（可下载权限），无权限字段直接跳过
            if is_permitted is not None and attr and not is_permitted(attr):
                continue

            # 列名未自定义（等于 attribute）时，优先取模型的 verbose_name 作为展示名
            if attr and column_name == attr:
                try:
                    model_field = resource._meta.model._meta.get_field(attr)
                    label = str(getattr(model_field, "verbose_name", attr))
                except FieldDoesNotExist:
                    label = str(attr)
            else:
                # column_name 已被自定义（如 dehydrate_full_name 设置了 column_name）
                label = str(column_name)

            result.append({
                "value": attr or column_name,  # 无 attribute 时回退到 column_name
                "label": label,
            })
        return success_response(message="导出字段列表获取成功", data=result)
